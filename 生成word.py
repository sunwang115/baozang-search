from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 设置默认字体
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)

# 标题
title = doc.add_heading('宝藏搜 - 搭建教程', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 副标题
subtitle = doc.add_paragraph('全能网盘搜索与自动化管理系统搭建指南')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.color.rgb = RGBColor(128, 128, 128)

doc.add_paragraph()

# 目录
doc.add_heading('目录', 1)
toc_items = [
    '环境要求',
    '安装步骤', 
    'Linux 一键部署',
    '常见问题',
    '功能与接口',
    'API 开放接口',
    '网盘配置',
    '联系方式',
    '免责声明'
]
for item in toc_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 环境要求
doc.add_heading('环境要求', 1)
doc.add_heading('系统要求', 2)

table = doc.add_table(rows=6, cols=2)
table.style = 'Light Grid Accent 1'
data = [
    ['项目', '要求'],
    ['操作系统', 'Windows 7+/Linux(Ubuntu/CentOS)/macOS'],
    ['Python', '3.8 及以上版本'],
    ['MySQL', '5.7 或 8.0+'],
    ['内存', '建议 2GB 以上'],
    ['硬盘', '建议 10GB 以上']
]
for i, (k, v) in enumerate(data):
    table.rows[i].cells[0].text = k
    table.rows[i].cells[1].text = v

doc.add_paragraph()
doc.add_heading('Python 版本检查', 2)
doc.add_paragraph('python --version')
doc.add_paragraph('# 或')
doc.add_paragraph('python3 --version')
doc.add_paragraph('如果版本低于 3.8，请先升级 Python。')

doc.add_heading('MySQL 版本检查', 2)
doc.add_paragraph('mysql --version')

doc.add_page_break()

# 安装步骤
doc.add_heading('安装步骤', 1)

doc.add_heading('第一步：下载源码', 2)
doc.add_paragraph('将项目源码上传到服务器：')
doc.add_paragraph('方式1：git 克隆')
doc.add_paragraph('git clone <你的仓库地址>')
doc.add_paragraph('cd baozang')
doc.add_paragraph('方式2：直接上传压缩包并解压')
doc.add_paragraph('上传后解压到任意目录')

doc.add_heading('第二步：安装依赖', 2)
doc.add_paragraph('Windows 系统：')
doc.add_paragraph('# 创建虚拟环境（推荐）')
doc.add_paragraph('python -m venv venv')
doc.add_paragraph('# 激活虚拟环境')
doc.add_paragraph('venv\\Scripts\\activate')
doc.add_paragraph('# 升级 pip')
doc.add_paragraph('pip install --upgrade pip')
doc.add_paragraph('# 安装依赖')
doc.add_paragraph('pip install -r requirements.txt')
doc.add_paragraph()
doc.add_paragraph('Linux/macOS 系统：')
doc.add_paragraph('# 创建虚拟环境')
doc.add_paragraph('python3 -m venv venv')
doc.add_paragraph('# 激活虚拟环境')
doc.add_paragraph('source venv/bin/activate')
doc.add_paragraph('# 升级 pip')
doc.add_paragraph('pip install --upgrade pip')
doc.add_paragraph('# 安装依赖')
doc.add_paragraph('pip install -r requirements.txt')

doc.add_heading('第三步：配置环境变量', 2)
doc.add_paragraph('将项目根目录的 .env.example 复制为 .env：')
doc.add_paragraph('cp .env.example .env')
doc.add_paragraph()
doc.add_paragraph('编辑 .env 文件，填写以下配置：')
doc.add_paragraph('# 系统密钥 (用于JWT签名，请设置复杂随机字符串)')
doc.add_paragraph('SECRET_KEY = your_random_secret_key_here')
doc.add_paragraph()
doc.add_paragraph('# MySQL 数据库配置')
doc.add_paragraph('DB_HOST = localhost')
doc.add_paragraph('DB_PORT = 3306')
doc.add_paragraph('DB_DATABASE = ucmao_search')
doc.add_paragraph('DB_USER = root')
doc.add_paragraph('DB_PASSWORD = your_mysql_password')
doc.add_paragraph('DB_CHARSET = utf8mb4')
doc.add_paragraph()
doc.add_paragraph('# 管理员账号配置')
doc.add_paragraph('ADMIN_USERNAME = admin')
doc.add_paragraph('ADMIN_PASSWORD = your_admin_password')

doc.add_heading('第四步：初始化数据库', 2)
doc.add_paragraph('方式1：手动初始化')
doc.add_paragraph('# 登录 MySQL')
doc.add_paragraph('mysql -u root -p')
doc.add_paragraph('# 创建数据库')
doc.add_paragraph('CREATE DATABASE IF NOT EXISTS ucmao_search DEFAULT CHARACTER SET utf8mb4;')
doc.add_paragraph('# 退出 MySQL')
doc.add_paragraph('exit')
doc.add_paragraph('# 导入表结构')
doc.add_paragraph('mysql -u root -p ucmao_search < schema.sql')
doc.add_paragraph()
doc.add_paragraph('方式2：使用脚本初始化')
doc.add_paragraph('python init_db.py')
doc.add_paragraph('注意：脚本会自动读取 .env 配置并创建数据库。')

doc.add_heading('第五步：启动服务', 2)
doc.add_paragraph('python app.py')
doc.add_paragraph()
doc.add_paragraph('看到以下输出表示启动成功：')
doc.add_paragraph(' * Running on http://127.0.0.1:5004')
doc.add_paragraph(' * Running on http://192.168.1.100:5004')

doc.add_heading('第六步：访问系统', 2)
doc.add_paragraph('打开浏览器访问：')
doc.add_paragraph('本机：http://127.0.0.1:5004')
doc.add_paragraph('局域网：http://你的IP:5004')

doc.add_heading('第七步：后台配置', 2)
doc.add_paragraph('1. 访问 http://127.0.0.1:5004/admin')
doc.add_paragraph('2. 使用 .env 中配置的管理员账号登录')
doc.add_paragraph('3. 进入"系统配置"页面')
doc.add_paragraph('4. 配置网盘凭证（Cookie/Token）')
doc.add_paragraph('5. 开启"公开聚合接口"（如需对外提供API）')

doc.add_page_break()

# Linux一键部署
doc.add_heading('Linux 一键部署', 1)
doc.add_paragraph('项目提供 install.sh 一键部署脚本，支持 Ubuntu/Debian/CentOS/Rocky/AlmaLinux：')
doc.add_paragraph('# 赋予执行权限')
doc.add_paragraph('chmod +x install.sh')
doc.add_paragraph('# 运行安装脚本')
doc.add_paragraph('sudo ./install.sh')
doc.add_paragraph()
doc.add_paragraph('脚本会自动完成：')
auto_items = [
    '安装 Python3、MySQL、Nginx',
    '创建虚拟环境',
    '安装项目依赖',
    '初始化数据库',
    '配置 Systemd 服务（开机自启）',
    '配置 Nginx 反向代理'
]
for item in auto_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 常见问题
doc.add_heading('常见问题', 1)

doc.add_heading('安装界面卡在第四步', 2)
doc.add_paragraph('问题描述：数据库初始化时卡住或报错。')
doc.add_paragraph()
doc.add_paragraph('解决方案：')
doc.add_paragraph('1. 检查 MySQL 服务是否启动：')
doc.add_paragraph('   Windows: net start mysql')
doc.add_paragraph('   Linux: systemctl status mysql')
doc.add_paragraph('2. 检查数据库连接信息是否正确（.env 文件）')
doc.add_paragraph('3. 手动创建数据库后重试：')
doc.add_paragraph('   CREATE DATABASE ucmao_search CHARACTER SET utf8mb4;')

doc.add_heading('报错 500', 2)
doc.add_paragraph('问题描述：页面显示 500 Internal Server Error。')
doc.add_paragraph()
doc.add_paragraph('解决方案：')
doc.add_paragraph('1. 查看日志文件 logs/search_ucmao.log')
doc.add_paragraph('2. 常见原因：')
doc.add_paragraph('   - 数据库连接失败 → 检查 .env 配置')
doc.add_paragraph('   - 缺少依赖 → 重新运行 pip install -r requirements.txt')
doc.add_paragraph('   - 端口被占用 → 修改 app.py 中的端口号')

doc.add_heading('转存中断', 2)
doc.add_paragraph('问题描述：网盘转存过程中断或失败。')
doc.add_paragraph()
doc.add_paragraph('解决方案：')
doc.add_paragraph('1. 检查网盘 Cookie/Token 是否过期')
doc.add_paragraph('2. 检查网盘容量是否已满')
doc.add_paragraph('3. 检查网络连接是否稳定')
doc.add_paragraph('4. 查看日志获取详细错误信息')

doc.add_heading('伪静态设置', 2)
doc.add_paragraph('如需使用域名访问并开启伪静态，Nginx 配置如下：')
doc.add_paragraph('server {')
doc.add_paragraph('    listen 80;')
doc.add_paragraph('    server_name your-domain.com;')
doc.add_paragraph('    ')
doc.add_paragraph('    location / {')
doc.add_paragraph('        proxy_pass http://127.0.0.1:5004;')
doc.add_paragraph('        proxy_set_header Host $host;')
doc.add_paragraph('        proxy_set_header X-Real-IP $remote_addr;')
doc.add_paragraph('        proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;')
doc.add_paragraph('    }')
doc.add_paragraph('    ')
doc.add_paragraph('    location /static {')
doc.add_paragraph('        alias /path/to/baozang/static;')
doc.add_paragraph('        expires 30d;')
doc.add_paragraph('    }')
doc.add_paragraph('}')

doc.add_heading('局域网无法访问', 2)
doc.add_paragraph('问题描述：其他设备无法访问 http://192.168.1.100:5004')
doc.add_paragraph()
doc.add_paragraph('解决方案：')
doc.add_paragraph('1. 检查防火墙是否放行 5004 端口：')
doc.add_paragraph('   Windows: netsh advfirewall firewall add rule name="宝藏搜" dir=in action=allow protocol=tcp localport=5004')
doc.add_paragraph('   Linux: sudo ufw allow 5004/tcp')
doc.add_paragraph('2. 检查网络类型是否为"专用网络"：')
doc.add_paragraph('   Get-NetConnectionProfile | Set-NetConnectionProfile -NetworkCategory Private')
doc.add_paragraph('3. 检查路由器是否开启 AP 隔离（关闭它）')

doc.add_heading('每日更新', 2)
doc.add_paragraph('系统内置定时任务，每 30 分钟自动清理过期临时分享。无需手动操作。')

doc.add_page_break()

# 功能与接口
doc.add_heading('功能与接口', 1)

doc.add_heading('全网搜', 2)
doc.add_paragraph('宝藏搜支持多网盘聚合搜索，包括：')
search_items = [
    '本地资源库搜索',
    '第三方 API 搜索（趣盘搜、狗狗盘搜等）',
    '外部数据源搜索'
]
for item in search_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('微信对话平台', 2)
doc.add_paragraph('可通过 API 接口对接微信公众号/小程序，实现微信内搜索资源。')

doc.add_heading('前端页面代码', 2)
doc.add_paragraph('前端采用响应式设计，支持：')
frontend_items = [
    'PC 端浏览器访问',
    '移动端浏览器访问',
    '资源封面展示',
    '分类浏览'
]
for item in frontend_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# API开放接口
doc.add_heading('API 开放接口', 1)

doc.add_heading('资源搜索接口', 2)
doc.add_paragraph('接口地址：GET /api')
doc.add_paragraph('请求参数：')

api_table = doc.add_table(rows=6, cols=4)
api_table.style = 'Light Grid Accent 1'
api_data = [
    ['参数', '类型', '必填', '说明'],
    ['keyword', 'string', '是', '搜索关键词'],
    ['cloud_name', 'string', '否', '网盘类型筛选'],
    ['type', 'string', '否', '资源类型'],
    ['limit', 'int', '否', '返回数量，默认100'],
    ['sort', 'string', '否', '排序方式：default/random/asc/desc']
]
for i, row_data in enumerate(api_data):
    for j, cell_data in enumerate(row_data):
        api_table.rows[i].cells[j].text = cell_data

doc.add_paragraph('返回示例：')
doc.add_paragraph('{')
doc.add_paragraph('  "success": true,')
doc.add_paragraph('  "total": 2,')
doc.add_paragraph('  "results": [')
doc.add_paragraph('    {')
doc.add_paragraph('      "source": "hot",')
doc.add_paragraph('      "name": "庆余年第二季",')
doc.add_paragraph('      "share_link": "https://pan.quark.cn/s/xxxx",')
doc.add_paragraph('      "cloud_name": "夸克网盘"')
doc.add_paragraph('    }')
doc.add_paragraph('  ]')
doc.add_paragraph('}')

doc.add_heading('资源详情接口', 2)
doc.add_paragraph('接口地址：POST /api/view-link')
doc.add_paragraph('请求参数：')

detail_table = doc.add_table(rows=4, cols=4)
detail_table.style = 'Light Grid Accent 1'
detail_data = [
    ['参数', '类型', '必填', '说明'],
    ['url', 'string', '是', '原始分享链接'],
    ['title', 'string', '否', '资源标题'],
    ['netdisk_name', 'string', '否', '网盘名称']
]
for i, row_data in enumerate(detail_data):
    for j, cell_data in enumerate(row_data):
        detail_table.rows[i].cells[j].text = cell_data

doc.add_paragraph('返回示例：')
doc.add_paragraph('{')
doc.add_paragraph('  "success": true,')
doc.add_paragraph('  "url": "https://pan.quark.cn/s/xxxx",')
doc.add_paragraph('  "mode": "copy"')
doc.add_paragraph('}')

doc.add_heading('网盘转存分享接口', 2)
doc.add_paragraph('接口地址：POST /create_share')
doc.add_paragraph('请求参数：')

share_table = doc.add_table(rows=4, cols=4)
share_table.style = 'Light Grid Accent 1'
share_data = [
    ['参数', '类型', '必填', '说明'],
    ['title', 'string', '是', '资源标题'],
    ['url', 'string', '是', '分享链接'],
    ['netdisk_name', 'string', '是', '网盘名称']
]
for i, row_data in enumerate(share_data):
    for j, cell_data in enumerate(row_data):
        share_table.rows[i].cells[j].text = cell_data

doc.add_paragraph('返回示例：')
doc.add_paragraph('{')
doc.add_paragraph('  "success": true,')
doc.add_paragraph('  "message": "分享创建成功"')
doc.add_paragraph('}')

doc.add_page_break()

# 网盘配置
doc.add_heading('网盘配置', 1)

doc.add_heading('夸克网盘', 2)
doc.add_paragraph('1. 登录夸克网盘网页版 https://pan.quark.cn')
doc.add_paragraph('2. 按 F12 打开开发者工具')
doc.add_paragraph('3. 刷新页面，在 Network 中找到第一个请求')
doc.add_paragraph('4. 复制 Request Headers 中的 Cookie')
doc.add_paragraph('5. 粘贴到后台"系统配置"中')

doc.add_heading('阿里网盘', 2)
doc.add_paragraph('1. 登录阿里云盘网页版 https://www.aliyundrive.com')
doc.add_paragraph('2. 使用开发者工具获取 refresh_token')
doc.add_paragraph('3. 粘贴到后台"系统配置"中')

doc.add_heading('百度网盘', 2)
doc.add_paragraph('1. 登录百度网盘网页版 https://pan.baidu.com')
doc.add_paragraph('2. 按 F12 打开开发者工具')
doc.add_paragraph('3. 刷新页面，复制 Cookie')
doc.add_paragraph('4. 粘贴到后台"系统配置"中')

doc.add_heading('UC网盘', 2)
doc.add_paragraph('1. 登录 UC网盘网页版 https://drive.uc.cn')
doc.add_paragraph('2. 按 F12 打开开发者工具')
doc.add_paragraph('3. 刷新页面，复制 Cookie')
doc.add_paragraph('4. 粘贴到后台"系统配置"中')

doc.add_heading('迅雷云盘', 2)
doc.add_paragraph('迅雷网盘需要三个参数：')
doc.add_paragraph('- Refresh Token')
doc.add_paragraph('- Captcha Sign')
doc.add_paragraph('- User ID')
doc.add_paragraph()
doc.add_paragraph('这三个参数需要同时填写，缺一不可。')

doc.add_page_break()

# 联系方式
doc.add_heading('联系方式', 1)
doc.add_paragraph('如有问题或建议，欢迎通过以下方式联系：')
doc.add_paragraph('项目地址：<你的仓库地址>', style='List Bullet')
doc.add_paragraph('问题反馈：请提交 Issue', style='List Bullet')

# 免责声明
doc.add_heading('免责声明', 1)
doc.add_paragraph('本工具仅供技术交流学习，严禁用于任何非法目的。因使用本工具造成的任何账号封禁或法律风险，均与原作者无关。')

# 结尾
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('宝藏搜 - 让每一份网盘资源都为你创造价值')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0, 102, 204)

# 保存
doc.save('宝藏搜-搭建教程.docx')
print('Word文档已生成：宝藏搜-搭建教程.docx')
